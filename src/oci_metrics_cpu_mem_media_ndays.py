import os
import time
import csv
import argparse
import json
import math
import sys
from pathlib import Path
from finops_period import validate_days, prompt_days, metric_interval, safe_start
from datetime import datetime, timedelta, timezone

import oci
from oci.monitoring.models import SummarizeMetricsDataDetails
from openpyxl import Workbook
from openpyxl.styles import PatternFill

# ================= CONFIGURAÇÕES =================
DAYS = 30
INTERVAL = "5m"

CPU_LOW = 5
CPU_MED = 15
CPU_HIGH = 80

MEM_LOW = 40
MEM_HIGH = 85

MAX_RETRIES = 3
RETRY_SLEEP = 3

homedir = os.path.expanduser("~")
CSV_PATH = os.path.join(homedir, f"Relatorio_CPU_Memoria_media_{DAYS}d_multi_region.csv")
XLSX_PATH = os.path.join(homedir, f"Relatorio_CPU_Memoria_media_{DAYS}d_multi_region.xlsx")
# ================================================

cfg = None
tenancy_id = None
identity = None
ISSUES = []
METADATA = {}

# ---------- helpers ----------
def get_regions():
    return [r.region_name for r in identity.list_region_subscriptions(tenancy_id).data]

def get_compartments():
    comps = oci.pagination.list_call_get_all_results(
        identity.list_compartments,
        tenancy_id,
        compartment_id_in_subtree=True
    ).data
    root = identity.get_compartment(tenancy_id).data
    return [c for c in comps if c.lifecycle_state == "ACTIVE"] + [root]

def mean_p95(values):
    if not values:
        return None, None
    values = sorted(values)
    mean = sum(values) / len(values)
    rank = .95 * (len(values) - 1)
    lo, hi = math.floor(rank), math.ceil(rank)
    p95 = values[lo] + (values[hi] - values[lo]) * (rank - lo)
    return mean, p95

def summarize_with_retry(monitoring, compartment_id, details):
    requested_start = details.start_time
    for attempt in range(1, MAX_RETRIES + 1):
        details.start_time = safe_start(requested_start, details.end_time, INTERVAL, datetime.now(timezone.utc))
        try:
            return monitoring.summarize_metrics_data(
                compartment_id=compartment_id,
                summarize_metrics_data_details=details
            )
        except oci.exceptions.ServiceError as e:
            if (e.status == 429 or e.status >= 500) and attempt < MAX_RETRIES:
                time.sleep(RETRY_SLEEP * attempt)
                continue
            raise

def get_metric(monitoring, compartment_id, instance_id, metric, start, end):
    query = f'{metric}[{INTERVAL}]{{resourceId = "{instance_id}"}}.mean()'
    details = SummarizeMetricsDataDetails(
        namespace="oci_computeagent",
        query=query,
        resolution=INTERVAL,
        start_time=start,
        end_time=end,
    )
    try:
        resp = summarize_with_retry(monitoring, compartment_id, details)
        if len(resp.data) > 1:
            raise ValueError("Mais de uma série retornada para a instância.")
    except Exception as exc:
        issue = f"{instance_id} / {metric}: {type(exc).__name__}: {exc}"
        ISSUES.append(issue)
        METADATA[metric] = {"start": "", "samples": 0, "error": issue}
        print(f"Falha na métrica {metric}; continuando a coleta.")
        return None, None
    METADATA[metric] = {"start": details.start_time.isoformat(), "samples": 0, "error": ""}
    if not resp.data or not resp.data[0].aggregated_datapoints:
        return None, None
    values = [d.value for d in resp.data[0].aggregated_datapoints if d.value is not None and math.isfinite(d.value)]
    METADATA[metric]["samples"] = len(values)
    return mean_p95(values)

def parse_baseline(instance):
    """
    Extrai baseline de OCPU corretamente via shape_config
    Compatível com SDK OCI atual
    """

    shape_cfg = getattr(instance, "shape_config", None)
    baseline = getattr(shape_cfg, "baseline_ocpu_utilization", None)

    if not baseline:
        return "NO", "Desativada", ""

    mapping = {
        "BASELINE_1_8": "12.5%",
        "BASELINE_1_2": "50%",
        "BASELINE_1_1": "100%"
    }

    enabled = "YES" if baseline in ("BASELINE_1_8", "BASELINE_1_2") else "NO"
    return enabled, mapping.get(baseline, baseline), baseline

def finops(cpu_mean, cpu_p95, mem_mean, mem_p95):
    if any(v is None for v in (cpu_mean, mem_mean, cpu_p95, mem_p95)):
        return "INSUFFICIENT_DATA"
    if cpu_p95 > CPU_HIGH or mem_p95 > MEM_HIGH:
        return "UPSCALE"

    if cpu_mean < CPU_LOW and mem_mean < MEM_LOW:
        return "DOWNSIZE-STRONG"
    if cpu_mean < CPU_MED and mem_mean < 60:
        return "DOWNSIZE"
    if mem_mean < MEM_LOW:
        return "DOWNSIZE-MEM"
    return "KEEP"

# ---------- main ----------
def main(days=30, outdir=None):
    global cfg, tenancy_id, identity, DAYS, INTERVAL, CSV_PATH, XLSX_PATH
    DAYS = validate_days(days)
    INTERVAL = metric_interval(DAYS)
    directory = Path(outdir or Path.home()).expanduser()
    directory.mkdir(parents=True, exist_ok=True)
    CSV_PATH = str(directory / f"Relatorio_CPU_Memoria_media_{DAYS}d_multi_region.csv")
    XLSX_PATH = str(Path(CSV_PATH).with_suffix(".xlsx"))
    ISSUES.clear()
    cfg = oci.config.from_file(profile_name=os.getenv("OCI_CLI_PROFILE", "DEFAULT"))
    tenancy_id = cfg["tenancy"]
    identity = oci.identity.IdentityClient(cfg)
    regions = get_regions()
    compartments = get_compartments()

    end = datetime.now(timezone.utc)
    start = end - timedelta(days=DAYS)

    rows = []
    print(f"Agregação: {INTERVAL}. No limite de retenção, o início tem margem de 5 minutos por consulta.")
    # Zera o CSV desta execução para impedir o reaproveitamento de resultados antigos.
    Path(CSV_PATH).write_text("", encoding="utf-8")

    print(f"\n📊 Coletando métricas dos últimos {DAYS} dias\n")

    for region in regions:
        print(f"\n🟢 Região: {region}")
        cfg_r = dict(cfg)
        cfg_r["region"] = region

        compute = oci.core.ComputeClient(cfg_r)
        monitoring = oci.monitoring.MonitoringClient(cfg_r)

        for comp in compartments:
            try:
                instances = oci.pagination.list_call_get_all_results(
                    compute.list_instances,
                    compartment_id=comp.id
                ).data
            except Exception as exc:
                issue = f"{region}/{comp.name}: {type(exc).__name__}: {exc}"
                ISSUES.append(issue)
                print(f"Falha de inventário: {issue}")
                continue

            running = [i for i in instances if i.lifecycle_state == "RUNNING"]
            if not running:
                continue

            print(f"  📁 {comp.name} | RUNNING: {len(running)}")

            for inst in running:
                detail_error = ""
                try:
                    inst_full = compute.get_instance(inst.id).data
                except Exception as exc:
                    inst_full = inst
                    detail_error = f"{inst.id}: {type(exc).__name__}: {exc}"
                    ISSUES.append(detail_error)

                cpu_mean, cpu_p95 = get_metric(
                    monitoring, comp.id, inst.id, "CpuUtilization", start, end
                )
                mem_mean, mem_p95 = get_metric(
                    monitoring, comp.id, inst.id, "MemoryUtilization", start, end
                )

                burst, baseline, baseline_raw = parse_baseline(inst_full)

                rows.append({
                    "region": region,
                    "compartment": comp.name,
                    "instance_name": inst.display_name,
                    "instance_ocid": inst.id,
                    "shape": inst.shape,
                    "ocpus": getattr(inst_full.shape_config, "ocpus", None),
                    "memory_gb": getattr(inst_full.shape_config, "memory_in_gbs", None),
                    "burstable_enabled": burst,
                    "baseline_percent": baseline,
                    "baseline_raw": baseline_raw,
                    "cpu_mean_percent": cpu_mean,
                    "cpu_p95_percent": cpu_p95,
                    "mem_mean_percent": mem_mean,
                    "mem_p95_percent": mem_p95,
                    "finops_recommendation": "INSUFFICIENT_DATA" if detail_error else finops(cpu_mean, cpu_p95, mem_mean, mem_p95),
                    "metric_interval": INTERVAL,
                    "requested_start_utc": start.isoformat(),
                    "end_utc": end.isoformat(),
                    "cpu_start_utc": METADATA["CpuUtilization"]["start"],
                    "mem_start_utc": METADATA["MemoryUtilization"]["start"],
                    "cpu_samples": METADATA["CpuUtilization"]["samples"],
                    "mem_samples": METADATA["MemoryUtilization"]["samples"],
                    "collection_error": " | ".join(filter(None, [detail_error, METADATA["CpuUtilization"]["error"], METADATA["MemoryUtilization"]["error"]]))
                })
                # Fecha o arquivo a cada recurso: progresso preservado mesmo se houver interrupção.
                with open(CSV_PATH, "a", newline="", encoding="utf-8") as checkpoint:
                    writer = csv.DictWriter(checkpoint, fieldnames=list(rows[-1]))
                    if len(rows) == 1:
                        writer.writeheader()
                    writer.writerow(rows[-1])

    headers = list(rows[0].keys()) if rows else ["instance_name", "finops_recommendation"]

    with open(CSV_PATH, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "FinOps"

    ws.append(headers)

    fill_keep = PatternFill("solid", fgColor="C6EFCE")
    fill_down = PatternFill("solid", fgColor="FFC7CE")
    fill_up = PatternFill("solid", fgColor="FFEB9C")

    rec_col = headers.index("finops_recommendation") + 1

    for r in rows:
        ws.append([r[h] for h in headers])
        row = ws.max_row
        rec = r["finops_recommendation"]
        if rec == "INSUFFICIENT_DATA":
            ws.cell(row=row, column=rec_col).fill = PatternFill("solid", fgColor="D9D9D9")
        elif rec.startswith("DOWNSIZE"):
            ws.cell(row=row, column=rec_col).fill = fill_down
        elif rec == "UPSCALE":
            ws.cell(row=row, column=rec_col).fill = fill_up
        else:
            ws.cell(row=row, column=rec_col).fill = fill_keep

    status = wb.create_sheet("Coleta")
    status.append(["Status", "PARCIAL" if ISSUES else "CONCLUÍDA"])
    status.append(["Instâncias", len(rows)])
    for issue in ISSUES:
        status.append(["Falha", issue])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(XLSX_PATH)
    status_path = directory / f"Relatorio_FinOps_Coleta_{DAYS}d.json"
    status_path.write_text(json.dumps({
        "status": "partial" if ISSUES else "complete", "days": DAYS, "interval": INTERVAL,
        "requested_start_utc": start.isoformat(), "end_utc": end.isoformat(),
        "instances": len(rows), "issues": ISSUES,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    from oci_metrics_cpu_mem_word_report import generate_report
    generate_report(rows=rows, days=DAYS, interval=INTERVAL, issues=ISSUES,
                    docx_path=directory / f"Relatorio_FinOps_CPU_Mem_{DAYS}d_multi_region.docx")
    print(f"Status: {'PARCIAL' if ISSUES else 'CONCLUÍDA'}; {len(rows)} instâncias, {len(ISSUES)} falhas.")
    print(f"Detalhes da coleta: {status_path}")

    print("\n✅ Relatórios gerados:")
    print(f"➡ CSV : {CSV_PATH}")
    print(f"➡ XLSX: {XLSX_PATH}")

    return 2 if ISSUES else 0


def cli():
    parser = argparse.ArgumentParser(description="Coleta OCI FinOps e gera CSV, Excel e Word.")
    parser.add_argument("--days", help="Dias de análise (1 a 90); omita para responder à pergunta.")
    parser.add_argument("--outdir", help="Destino dos relatórios (padrão: diretório pessoal).")
    args = parser.parse_args()
    try:
        value = args.days
        if value is None and not sys.stdin.isatty():
            value = os.getenv("METRICS_DAYS")
        days = prompt_days() if value is None else validate_days(value)
    except ValueError as exc:
        parser.error(str(exc))
    return main(days, args.outdir)


if __name__ == "__main__":
    raise SystemExit(cli())
