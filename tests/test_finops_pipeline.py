import contextlib
import csv
from datetime import datetime, timedelta, timezone
import io
import json
from pathlib import Path
import sys
import tempfile
from types import SimpleNamespace as NS
import unittest
from unittest.mock import Mock, patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
import oci_metrics_cpu_mem_media_ndays as collector
from finops_period import validate_days, metric_interval, prompt_days, safe_start
from docx import Document
from openpyxl import load_workbook


class PeriodTests(unittest.TestCase):
    def test_limits_and_resolution(self):
        for days, interval in [(1, "5m"), (29, "5m"), (30, "5m"), (31, "1h"), (90, "1h")]:
            self.assertEqual(metric_interval(days), interval)
        for invalid in [0, 91, -1, "abc", "1.5", 1.5, "", None]:
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                validate_days(invalid)

    def test_prompt_retries_and_default(self):
        with patch("builtins.input", side_effect=["91", "abc", ""]), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(prompt_days(), 30)

    def test_window_ages_during_collection(self):
        end = datetime(2026, 9, 9, 18, 50, 7, tzinfo=timezone.utc)
        for days in [30, 90]:
            start = end - timedelta(days=days)
            now = end + timedelta(hours=2)
            effective = safe_start(start, end, metric_interval(days), now)
            self.assertEqual(effective, now - timedelta(days=days) + timedelta(minutes=5))
        start = end - timedelta(days=29)
        self.assertEqual(safe_start(start, end, "5m", end), start)

    def test_retry_rechecks_clock(self):
        end = datetime.now(timezone.utc)
        details = collector.SummarizeMetricsDataDetails(
            namespace="oci_computeagent", query="CpuUtilization[5m].mean()",
            start_time=end - timedelta(days=30), end_time=end)
        client = Mock()
        starts = []
        def request(**kwargs):
            starts.append(kwargs["summarize_metrics_data_details"].start_time)
            if len(starts) == 1:
                raise collector.oci.exceptions.ServiceError(429, "TooManyRequests", {}, "retry")
            return NS(data=[])
        client.summarize_metrics_data.side_effect = request
        with patch.object(collector, "INTERVAL", "5m"), patch.object(collector, "datetime") as clock, patch.object(collector.time, "sleep"):
            clock.now.side_effect = [end, end + timedelta(minutes=10)]
            collector.summarize_with_retry(client, "comp", details)
        self.assertEqual(starts[1] - starts[0], timedelta(minutes=10))

    def test_missing_metrics_and_peaks(self):
        self.assertEqual(collector.finops(None, None, None, None), "INSUFFICIENT_DATA")
        self.assertEqual(collector.finops(4, 90, 20, 95), "UPSCALE")
        self.assertEqual(collector.mean_p95([0, 100]), (50, 95))

    def test_interactive_prompt_overrides_old_environment(self):
        with patch.object(sys, "argv", ["collector"]), patch.object(sys.stdin, "isatty", return_value=True), patch.dict("os.environ", {"METRICS_DAYS": "29"}), patch.object(collector, "prompt_days", return_value=90) as prompt, patch.object(collector, "main", return_value=0) as main:
            self.assertEqual(collector.cli(), 0)
            prompt.assert_called_once()
            main.assert_called_once_with(90, None, None, 90, 75)


class PipelineTests(unittest.TestCase):
    def run_pipeline(self, directory, days, mode):
        instance = NS(id="synthetic-instance", display_name="vm-test", lifecycle_state="RUNNING",
                      shape="VM.Standard.E4.Flex", shape_config=NS(ocpus=4, memory_in_gbs=16, baseline_ocpu_utilization="BASELINE_1_1"))
        compute = Mock()
        compute.get_instance.return_value = NS(data=instance)
        def paginate(method, **kwargs):
            if mode == "denied":
                raise RuntimeError("synthetic permission failure")
            return NS(data=[] if mode == "empty" else [instance])
        monitoring = Mock()
        def metric(**kwargs):
            details = kwargs["summarize_metrics_data_details"]
            self.assertEqual(details.resolution, metric_interval(days))
            retention = 30 if days <= 30 else 90
            self.assertGreater(details.start_time, datetime.now(timezone.utc) - timedelta(days=retention))
            if "Memory" in details.query and mode == "partial":
                raise collector.oci.exceptions.ServiceError(400, "InvalidParameter", {}, "synthetic metric failure")
            return NS(data=[NS(aggregated_datapoints=[NS(value=2), NS(value=3)])])
        monitoring.summarize_metrics_data.side_effect = metric
        with patch.object(collector.oci.config, "from_file", return_value={"tenancy": "synthetic"}), patch.object(collector.oci.identity, "IdentityClient"), patch.object(collector, "get_regions", return_value=["region-test"]), patch.object(collector, "get_compartments", return_value=[NS(id="comp", name="test")]), patch.object(collector.oci.core, "ComputeClient", return_value=compute), patch.object(collector.oci.monitoring, "MonitoringClient", return_value=monitoring), patch.object(collector.oci.pagination, "list_call_get_all_results", side_effect=paginate), contextlib.redirect_stdout(io.StringIO()):
            return collector.main(days, directory)

    def test_reports_complete_partial_empty_and_denied(self):
        for mode, days in [("complete", 30), ("complete", 90), ("partial", 30), ("empty", 7), ("denied", 31)]:
            with self.subTest(mode=mode, days=days), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                code = self.run_pipeline(tmp, days, mode)
                self.assertEqual(code, 2 if mode in ("partial", "denied") else 0)
                status = json.loads(next(root.glob("*.json")).read_text(encoding="utf-8"))
                self.assertEqual(status["status"], "partial" if code else "complete")
                with next(root.glob("*.csv")).open(encoding="utf-8", newline="") as stream:
                    rows = list(csv.DictReader(stream))
                self.assertEqual(len(rows), 0 if mode in ("empty", "denied") else 1)
                document = Document(next(root.glob("*.docx")))
                text = "\n".join(p.text for p in document.paragraphs)
                self.assertIn(f"{days} dias", text)
                self.assertIn(metric_interval(days), text)
                if mode == "partial":
                    self.assertEqual(rows[0]["finops_recommendation"], "INSUFFICIENT_DATA")
                    self.assertIn("PARCIAL", text)
                    self.assertIn("synthetic metric failure", text)
                if mode == "complete":
                    self.assertIn("vm-test", text)
                    self.assertEqual(rows[0]["burstable_enabled"], "NO")
                wb = load_workbook(next(root.glob("*.xlsx")))
                self.assertIn("Coleta", wb.sheetnames)
                wb.close()

    def test_empty_run_does_not_reuse_previous_csv(self):
        with tempfile.TemporaryDirectory() as tmp:
            self.run_pipeline(tmp, 30, "complete")
            self.run_pipeline(tmp, 30, "empty")
            doc = Document(next(Path(tmp).glob("*.docx")))
            self.assertNotIn("vm-test", "\n".join(p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
